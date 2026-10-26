"""
Document Parser Service.
Extracts clean, readable text from uploaded PDF and DOCX resume documents
with robust error handling for corrupted, encrypted, or empty files.
"""

import io
import logging
from typing import BinaryIO, Union

logger = logging.getLogger("DocumentParser")


class DocumentParsingError(Exception):
    """Raised when parsing a resume document fails."""
    pass


class UnsupportedFileTypeError(Exception):
    """Raised when an uploaded file is not supported."""
    pass


def extract_text_from_pdf(file_source: Union[BinaryIO, bytes]) -> str:
    """
    Extract text content from a PDF file stream or bytes.
    Supports pypdf and PyPDF2 transparently.
    """
    stream = io.BytesIO(file_source) if isinstance(file_source, bytes) else file_source
    text_chunks = []

    try:
        # Prefer modern pypdf, fallback to PyPDF2
        try:
            import pypdf
            reader = pypdf.PdfReader(stream)
        except ImportError:
            import PyPDF2
            reader = PyPDF2.PdfReader(stream)

        if reader.is_encrypted:
            # Attempt empty password decryption
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentParsingError("This PDF is password-protected. Please upload an unprotected version.")

        if len(reader.pages) == 0:
            raise DocumentParsingError("The uploaded PDF has no readable pages.")

        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_chunks.append(page_text.strip())

    except DocumentParsingError:
        raise
    except Exception as exc:
        logger.error(f"Failed to extract PDF text: {exc}")
        raise DocumentParsingError(f"Could not read PDF file: {str(exc)}")

    full_text = "\n\n".join(text_chunks).strip()
    if not full_text:
        raise DocumentParsingError(
            "No extractable text was found in the PDF. The file might contain scanned images without OCR. "
            "Please upload a text-based PDF or DOCX file."
        )

    return full_text


def extract_text_from_docx(file_source: Union[BinaryIO, bytes]) -> str:
    """
    Extract text content from a DOCX file stream or bytes.
    Extracts from paragraphs and tables.
    """
    stream = io.BytesIO(file_source) if isinstance(file_source, bytes) else file_source

    try:
        import docx
        doc = docx.Document(stream)
        
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract text from any tables inside the resume
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_texts.append(" | ".join(row_cells))

        combined = paragraphs + table_texts
        full_text = "\n".join(combined).strip()

    except Exception as exc:
        logger.error(f"Failed to extract DOCX text: {exc}")
        raise DocumentParsingError(f"Could not read DOCX document: {str(exc)}")

    if not full_text:
        raise DocumentParsingError("The uploaded DOCX document contains no readable text.")

    return full_text


def parse_resume_document(filename: str, file_bytes: bytes) -> str:
    """
    Main entry point for extracting text from an uploaded resume file.
    Validates file type and routes to the appropriate parser.
    """
    if not filename:
        raise UnsupportedFileTypeError("Filename cannot be empty.")

    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file format '{filename}'. Only .pdf and .docx files are supported."
        )
